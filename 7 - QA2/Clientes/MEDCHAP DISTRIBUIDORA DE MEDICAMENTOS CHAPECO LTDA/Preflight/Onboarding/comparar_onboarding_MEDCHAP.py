#!/usr/bin/env python3
"""Conferencia read-only do onboarding MEDCHAP (arquivos do "Baixar Tudo") x base Oracle (consultas O01-O03).
Sem DML/DDL. Agencia/conta so comparadas por SHA-256 dos digitos; saida mascarada."""
import collections, datetime as dt, hashlib, json, re, subprocess, sys, unicodedata, warnings
from pathlib import Path
import openpyxl
warnings.filterwarnings("ignore")
B = Path(__file__).resolve().parents[2]
ARQ = B / "Artefatos_Revisao" / "Onboarding_Arquivos"
ZIP = B / "Artefatos_Revisao" / "Onboarding - MEDCHAP DISTRIBUIDORA DE MEDICAMENTOS CHAPECO LTDA.zip"
LOG = B / "Logs"
dig = lambda s: re.sub(r"\D", "", str(s or ""))
h = lambda s: hashlib.sha256((s or "-").encode()).hexdigest().upper()
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()
def ora(name):
    d = json.load(open(LOG / f"{name}_20260925.json")); assert d["status"] == "CONSULTA_READ_ONLY_CONCLUIDA"; return d
res = {"base": "MEDCHAP DISTRIBUIDORA DE MEDICAMENTOS CHAPECO LTDA", "executado_em": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
       "fonte_onboarding": {"arquivo": ZIP.name, "origem": "Onboarding Deploy - botao Baixar Tudo (25/09/2026 09:30)", "api": "indisponivel (HTTP 403, aguardando Adelcione)"},
       "dml_executado": False, "ddl_executado": False}
o1, o2, o3 = ora("O01_Empresas"), ora("O02_Usuarios_Grupos"), ora("O03_Contas_Bancarias_Mascarado")
res["identidade_oracle"] = {k: o1[k] for k in ("session_user", "current_schema", "service_name", "db_name")}
# ---------- empresas ----------
txt = subprocess.run(["pdftotext", "-layout", str(ARQ / "01 - Empresas/Resumo - 01 - Dados da Empresa.pdf"), "-"], capture_output=True, text=True).stdout
emp_ob = []
for bloco in re.split(r"\n(?=(?:Matriz|Empresa \d+) — )", txt)[1:]:
    g = lambda k: (re.search(k + r":\s*(.+)", bloco) or [None, ""])[1].strip()
    emp_ob.append({"nome": bloco.split("—", 1)[1].split("\n")[0].strip(), "cnpj": dig(g("CNPJ/CPF")), "ie": g("IE"), "uf": g("UF"), "tipo": g("Tipo")})
emp_or = {r["CNPJ"]: r for r in o1["linhas"]}
cmp_emp = []
for e in emp_ob:
    r = emp_or.get(e["cnpj"])
    item = {"cnpj": e["cnpj"], "razao_onboarding": e["nome"], "tipo_onboarding": e["tipo"], "codemp": r and r["CODEMP"], "razao_oracle": r and r["RAZAOSOCIAL"]}
    if not r: item["status"] = "AUSENTE_NO_ORACLE"
    else:
        dif = []
        if dig(e["ie"]) != dig(r["IE"]) and norm(e["ie"]) != norm(r["IE"]): dif.append(f"IE onboarding={e['ie']} oracle={r['IE']}")
        if e["uf"] not in ("", "—") and e["uf"] != r["UF"]: dif.append(f"UF onboarding={e['uf']} oracle={r['UF']}")
        if e["uf"] in ("", "—"): item["observacao"] = f"UF nao informada no onboarding; Oracle={r['UF']}"
        if not norm(r["RAZAOSOCIAL"]) or not norm(e["nome"]).startswith(norm(r["RAZAOSOCIAL"])[:30]): dif.append("razao social diferente")
        if e["tipo"].upper() == "FILIAL" and e["cnpj"][:8] != "00577604": dif.append("marcada como FILIAL no onboarding, mas CNPJ de outra raiz (empresa distinta)")
        item["status"] = "DIVERGENTE" if dif else "CORRESPONDE"; item["divergencias"] = dif
    cmp_emp.append(item)
extras = [{"codemp": r["CODEMP"], "cnpj": c, "razao": r["RAZAOSOCIAL"]} for c, r in emp_or.items() if c not in {e["cnpj"] for e in emp_ob}]
res["empresas"] = {"onboarding": len(emp_ob), "oracle_tsiemp": len(emp_or), "comparacao": cmp_emp, "adicionais_no_oracle": extras,
                   "observacao": "quantidade de empresas nao e gate (BASE_INTEIRA)"}
# ---------- usuarios ----------
ws = openpyxl.load_workbook(ARQ / "08 - Planilhas/Usuarios_Sankhya_Medchap.xlsx", data_only=True)["Usuarios"]
rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)][1:]
EMP_MAP = {"MATRIZ": 1, "FILIAL - CBS": 2, "FILIAL - SJP": 3}  # CBS=Curitibanos(CODEMP 2), SJP=Sao Jose dos Pinhais(CODEMP 3) conforme TSIEMP/TSICID
ob = [{"linha": i + 2, "email": (r[1] or "").strip().lower(), "nome": norm(r[2]), "empresa": r[3], "codemp_esperado": EMP_MAP.get(r[3]),
       "codgrupo": r[4], "grupo": r[5]} for i, r in enumerate(rows)]
todos = o2["linhas"]; comg = [u for u in todos if (u["CODGRUPO"] or 0) > 0]
res_u = {"planilha_linhas": len(ob), "oracle_total": len(todos), "oracle_com_grupo": len(comg), "oracle_sem_grupo_modelo": len(todos) - len(comg),
         "usuarios_modelo_sem_grupo": [{"codusu": u["CODUSU"], "nome": u["NOMEUSU"]} for u in todos if not (u["CODGRUPO"] or 0) > 0],
         "criterio": "chave e-mail + NOMEUSU (a planilha repete e-mails compartilhados); somente TSIUSU.CODGRUPO>0", "mapa_empresa": EMP_MAP}
by_key = {(u["EMAIL"] or "", norm(u["NOMEUSU"])): u for u in comg}
by_nome = collections.defaultdict(list); [by_nome[norm(u["NOMEUSU"])].append(u) for u in comg]
usados, corr, div, aus = set(), [], [], []
for p in ob:
    u = by_key.get((p["email"], p["nome"]))
    via = "email+nome"
    if not u and len(by_nome.get(p["nome"], [])) == 1: u, via = by_nome[p["nome"]][0], "nome (e-mail diferente)"
    if not u:
        cand = [x for x in comg if x["EMAIL"] == p["email"]]
        if len(cand) == 1: u, via = cand[0], "email (nome diferente)"
    if not u: aus.append(p); continue
    usados.add(u["CODUSU"]); d = []
    if via != "email+nome": d.append(f"correspondencia por {via}: oracle {u['NOMEUSU']} <{u['EMAIL']}>")
    if p["codgrupo"] not in (None, "") and int(p["codgrupo"]) != u["CODGRUPO"]: d.append(f"CODGRUPO planilha={p['codgrupo']} oracle={u['CODGRUPO']}")
    if p["codemp_esperado"] and p["codemp_esperado"] != u["CODEMP"]: d.append(f"empresa planilha={p['empresa']}(CODEMP {p['codemp_esperado']}) oracle CODEMP={u['CODEMP']}")
    item = {"linha": p["linha"], "email": p["email"], "nome": p["nome"], "grupo_planilha": p["grupo"], "codgrupo_oracle": u["CODGRUPO"], "grupo_oracle": u["NOMEGRUPO"], "codusu": u["CODUSU"], "codemp_oracle": u["CODEMP"]}
    (div if d else corr).append({**item, **({"divergencias": d} if d else {})})
extra = [{"codusu": u["CODUSU"], "nome": u["NOMEUSU"], "email": u["EMAIL"], "codgrupo": u["CODGRUPO"], "codemp": u["CODEMP"]} for u in comg if u["CODUSU"] not in usados]
xw = collections.defaultdict(collections.Counter)
for x in corr + div: xw[x["grupo_planilha"]][f"{x['codgrupo_oracle']}-{x['grupo_oracle']}"] += 1
res_u.update({"correspondem": len(corr), "divergentes": div, "ausentes_no_oracle": aus, "extras_no_oracle": extra,
              "equivalencia_grupos_planilha_x_oracle": {k: dict(v) for k, v in sorted(xw.items())},
              "grupos_com_mais_de_um_destino": {k: dict(v) for k, v in xw.items() if len(v) > 1}})
res["usuarios"] = res_u
# ---------- contas bancarias ----------
ws = openpyxl.load_workbook(ARQ / "Contas Bancárias/MODELO SANKHYA Contas Bancarias.xlsx", data_only=True)["Contas"]
lab = {str(r[0]).strip(): r[1:] for r in ws.iter_rows(values_only=True) if r[0]}
cts = []
for i, cnpj in enumerate(lab.get("Empresa CNPJ:", [])):
    if not cnpj: continue
    g = lambda k: (lab.get(k) or [None] * 99)[i]
    banco = str(g("Banco:") or ""); cod = (re.search(r"CÓD\. (\d+)", banco) or [None, ""])[1]
    ag, ct = dig(g("Agência:")), dig(g("Conta:")); dv = dig(g("Dígito:"))
    dv_ag = str(g("Digito:") or "").strip().upper().replace("X", "0")  # Banco do Brasil: DV "X" gravado como "0"
    ag_hashes = {h(ag), h(ag + dig(dv_ag))} if dig(dv_ag) else {h(ag)}
    cts.append({"conta": i + 1, "cnpj": dig(cnpj), "codbco": int(cod) if cod else None, "banco": banco.split(" - ")[0], "ag_hash": h(ag), "ag_hashes": ag_hashes,
                "conta_hashes": {h(ct), h(ct + dv)}, "conta_final2": (ct + dv)[-2:], "emite_boleto": g("A conta emite boleto? (S/N)")})
tsicta = o3["linhas"]; cmp_c = []
for c in cts:
    m = [t for t in tsicta if t["CODBCO"] == c["codbco"] and (t["CONTA_HASH"] in c["conta_hashes"] or t["CONTA_SEM_DV_HASH"] in c["conta_hashes"])]
    it = {"conta_planilha": c["conta"], "cnpj": c["cnpj"], "codbco": c["codbco"], "banco": c["banco"], "conta_final2": c["conta_final2"]}
    if not m: it["status"] = "AUSENTE_NO_ORACLE"
    else:
        t = m[0]; d = []
        if t["CNPJ_EMP"] != c["cnpj"]: d.append(f"empresa oracle CODEMP={t['CODEMP']}")
        if t["AG_HASH"] not in c["ag_hashes"]: d.append("agencia diferente (comparacao por hash)")
        if t["ATIVA"] != "S": d.append(f"ATIVA={t['ATIVA']}")
        it.update({"codctabcoint": t["CODCTABCOINT"], "codemp": t["CODEMP"], "status": "DIVERGENTE" if d else "CORRESPONDE", "divergencias": d})
    cmp_c.append(it)
usadas = {x.get("codctabcoint") for x in cmp_c}
res["contas_bancarias"] = {"planilha": len(cts), "oracle_tsicta": len(tsicta), "comparacao": cmp_c,
    "extras_no_oracle": [{"codctabcoint": t["CODCTABCOINT"], "codemp": t["CODEMP"], "codbco": t["CODBCO"], "banco": t["NOMEBCO"], "ativa": t["ATIVA"], "conta_final2": t["CONTA_FINAL2"]} for t in tsicta if t["CODCTABCOINT"] not in usadas],
    "mascaramento": "agencia/conta comparadas por SHA-256 dos digitos; exibidos so os 2 ultimos digitos"}
out = B / "Onboarding_Comparacao_ReadOnly_MEDCHAP.json"
if out.exists(): sys.exit(f"saida existe: {out}")
out.write_text(json.dumps(res, ensure_ascii=False, indent=2, default=list) + "\n", encoding="utf-8")
print(json.dumps({"empresas": [(x["cnpj"], x["status"], x.get("divergencias"), x.get("observacao")) for x in cmp_emp], "empresas_extras": extras,
    "usuarios": {k: (len(v) if isinstance(v, list) else v) for k, v in res_u.items() if k not in ("usuarios_modelo_sem_grupo", "equivalencia_grupos_planilha_x_oracle", "mapa_empresa", "criterio")},
    "contas": [(x["conta_planilha"], x["banco"], x["status"], x.get("divergencias")) for x in cmp_c], "contas_extras": len(res["contas_bancarias"]["extras_no_oracle"])}, ensure_ascii=False, indent=1, default=list))
